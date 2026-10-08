"""Confronto con il foglio di riferimento (CALCOLO CONTO ESTERO.numbers).

Livello A (fedeltà formule): stessi cambi dell'originale -> tutte le colonne devono coincidere.
Livello B (cambi): cambi ufficiali Banca d'Italia vs cambi scritti a mano nell'originale.

Il ricalcolo usa Numbers (macOS): è lo stesso programma dell'utente. Run: python3 test_compare.py
"""
import datetime as dt
import os
import subprocess
import sys
import tempfile

import openpyxl
import pandas as pd

import core

HERE = os.path.dirname(os.path.abspath(__file__))
ORIGINALE = os.path.join(HERE, "CALCOLO CONTO ESTERO.numbers")
COLONNE = {"E": "saldo", "F": "gg", "G": "saldo x gg", "J": "entrate EUR", "K": "uscite EUR",
           "L": "saldo EUR", "M": "gg EUR", "N": "saldo x gg EUR"}
TOL = 1e-6


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, dt.timedelta):
        return float(v.days)
    return float(v)


def load_original():
    from numbers_parser import Document  # valori già calcolati dal file Numbers
    return [list(r) for r in Document(ORIGINALE).sheets[0].tables[0].rows(values_only=True)[1:]]


def recalc(xlsx_bytes: bytes):
    """Apre il file in Numbers, lo riesporta con i valori calcolati e restituisce il foglio Movimenti."""
    with tempfile.TemporaryDirectory() as d:
        src, dst = os.path.join(d, "in.xlsx"), os.path.join(d, "out.xlsx")
        open(src, "wb").write(xlsx_bytes)
        script = (f'tell application "Numbers"\nset d to open POSIX file "{src}"\ndelay 3\n'
                  f'export d to POSIX file "{dst}" as Microsoft Excel\nclose d saving no\nend tell')
        subprocess.run(["osascript", "-e", script], check=True, capture_output=True, timeout=120)
        return openpyxl.load_workbook(dst, data_only=True)["Movimenti"]


def movimenti_da_originale(rows):
    return pd.DataFrame({"data": [pd.Timestamp(r[0]) for r in rows], "causale": [r[1] or "" for r in rows],
                         "entrate": [r[2] or 0.0 for r in rows], "uscite": [r[3] or 0.0 for r in rows],
                         "saldo_dichiarato": [r[4] for r in rows]})


def test_livello_A_formule_identiche_all_originale():
    rows = load_original()
    mov = movimenti_da_originale(rows)
    rates = pd.DataFrame({"data": mov["data"], "cambio": [r[8] for r in rows]}).drop_duplicates("data")
    opening = (2024, 12)
    xlsx = core.build_workbook(mov, rates, "GBP", opening, {opening: rows[0][8]})
    ws = recalc(xlsx)
    errori = []
    for i, r in enumerate(rows):
        for col, nome in COLONNE.items():
            got, want = num(ws[f"{col}{i + 2}"].value), num(r[openpyxl.utils.column_index_from_string(col) - 1])
            if abs(got - want) > TOL:
                errori.append((i + 2, nome, got, want))
    assert not errori, f"{len(errori)} celle diverse, prime 8: {errori[:8]}"


def test_livello_B_cambi_originali_sono_quelli_di_7_giorni_prima():
    """Caratterizza l'originale: ogni cambio = ultimo cambio ufficiale disponibile al (data - 7 giorni)."""
    rows = load_original()[1:]  # la riga 1 è il saldo iniziale
    lo, hi = rows[0][0].date() - dt.timedelta(days=20), rows[-1][0].date()
    g = core.daily_rates("GBP", lo, hi)
    uguali = sum(abs(core.asof_rate(g, pd.Timestamp(r[0]) - pd.Timedelta(days=7))[1] - r[8]) < 1e-9 for r in rows)
    stesso_giorno = sum(abs(core.asof_rate(g, pd.Timestamp(r[0]))[1] - r[8]) < 1e-9 for r in rows)
    print(f"  cambi originali = ufficiale a data-7gg: {uguali}/{len(rows)}; = ufficiale del giorno: {stesso_giorno}/{len(rows)}")
    # eccezioni note: 20/04/2025 e 03/05/2025 (x2) usano un giorno lavorativo ancora prima (Pasqua / 25 aprile)
    assert uguali >= len(rows) - 3


if __name__ == "__main__":
    fallite = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            try:
                fn()
                print("ok  ", name)
            except Exception as exc:  # noqa: BLE001 - riepilogo leggibile
                fallite += 1
                print("FAIL", name, "->", str(exc)[:600])
    sys.exit(1 if fallite else 0)
