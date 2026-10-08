"""Cambi Banca d'Italia (ex UIC) + costruzione del foglio movimenti in euro."""
from __future__ import annotations

import datetime as dt
import io
import os
import tempfile
from dataclasses import dataclass

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE = "https://tassidicambio.bancaditalia.it/terzevalute-wf-web/rest/v1.0"
HEADERS = {"Accept": "application/json"}
TIMEOUT = 30
LOOKBACK_DAYS = 10  # copre weekend + festivi consecutivi prima del primo movimento
SOURCE = "Banca d'Italia - tassi di cambio (ex UIC)"
SOURCE_MONTHLY = "Medie mensili Banca d'Italia = provvedimenti Agenzia delle Entrate"
SALDO_TOLERANCE = 0.01


class RatesError(RuntimeError):
    pass


def _get(path: str, **params) -> dict:
    try:
        r = requests.get(f"{BASE}/{path}", params=params, headers=HEADERS, timeout=TIMEOUT)
        r.raise_for_status()
        return r.json()
    except (requests.RequestException, ValueError) as exc:
        raise RatesError(f"Banca d'Italia non raggiungibile o risposta non valida: {exc}") from exc


def _to_units_per_eur(rate: float, convention: str) -> float:
    """Normalizza a 'unità di valuta estera per 1 EUR' (convenzione usata nel foglio)."""
    return 1 / rate if "euro per 1" in convention.lower() else rate


def list_currencies() -> dict[str, str]:
    """{codice ISO: nome} delle valute ancora in vigore."""
    out = {}
    for c in _get("currencies", lang="it")["currencies"]:
        if any(x.get("validityEndDate") is None for x in c["countries"]) and c["isoCode"] != "EUR":
            out[c["isoCode"]] = c["name"]
    return dict(sorted(out.items()))


def daily_rates(iso: str, start: dt.date, end: dt.date) -> pd.DataFrame:
    """Cambi giornalieri: colonne data, cambio (valuta per 1 EUR). Solo giorni con quotazione."""
    data = _get("dailyTimeSeries", startDate=start.isoformat(), endDate=end.isoformat(),
                baseCurrencyIsoCode=iso, currencyIsoCode="EUR", lang="it")
    rows = [(pd.Timestamp(r["referenceDate"]),
             _to_units_per_eur(float(r["avgRate"]), r["exchangeConvention"]))
            for r in data.get("rates", [])]
    if not rows:
        raise RatesError(f"Nessun cambio {iso} tra {start} e {end}.")
    return pd.DataFrame(rows, columns=["data", "cambio"]).sort_values("data").reset_index(drop=True)


def monthly_average(iso: str, year: int, month: int) -> float | None:
    """Cambio medio mensile (stessi valori dei provvedimenti dell'Agenzia delle Entrate)."""
    data = _get("monthlyAverageRates", month=month, year=year,
                baseCurrencyIsoCode=iso, currencyIsoCode="EUR", lang="it")
    rates = data.get("rates") or []
    if not rates:
        return None
    return _to_units_per_eur(float(rates[0]["avgRate"]), rates[0]["exchangeConvention"])


def asof_rate(rates: pd.DataFrame, day: pd.Timestamp) -> tuple[pd.Timestamp, float]:
    """Ultimo cambio con data <= day (weekend/festivi -> giorno precedente)."""
    past = rates[rates["data"] <= day]
    if past.empty:
        raise RatesError(f"Nessun cambio disponibile fino al {day:%d/%m/%Y}.")
    last = past.iloc[-1]
    return last["data"], float(last["cambio"])


@dataclass(frozen=True)
class Mapping:
    data: str
    causale: str | None
    entrate: str | None
    uscite: str | None
    importo: str | None  # alternativa a entrate/uscite: importo con segno
    saldo: str | None


def guess_mapping(columns: list[str]) -> Mapping:
    def find(*keys: str) -> str | None:
        return next((c for c in columns if any(k in str(c).lower() for k in keys)), None)
    split = find("entrat", "avere", "accredit")
    return Mapping(data=find("data", "giorno", "date") or columns[0],
                   causale=find("causale", "descr"),
                   entrate=split, uscite=find("uscit", "dare", "addebit"),
                   importo=None if split else find("importo", "movimento"),
                   saldo=find("saldo"))


def read_table(name: str, content: bytes) -> pd.DataFrame:
    """Legge il primo foglio di un file .xlsx / .ods / .numbers."""
    ext = name.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm"):
        df = pd.read_excel(io.BytesIO(content), engine="openpyxl")
    elif ext == "ods":
        df = pd.read_excel(io.BytesIO(content), engine="odf")
    elif ext == "numbers":
        from numbers_parser import Document
        with tempfile.NamedTemporaryFile(suffix=".numbers", delete=False) as f:
            f.write(content)
        try:
            rows = Document(f.name).sheets[0].tables[0].rows(values_only=True)
        finally:
            os.unlink(f.name)
        df = pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]])
    else:
        raise ValueError(f"Formato non supportato: .{ext} (usa xlsx, ods o numbers)")
    df = df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    df.columns = unique_names(df.columns)
    return df


def unique_names(names) -> list[str]:
    """Nomi colonna univoci: vuoti -> 'colonna N', duplicati -> 'saldo', 'saldo.1', ..."""
    seen: dict[str, int] = {}
    out = []
    for i, n in enumerate(names):
        base = str(n).strip()
        if base in ("", "None", "nan") or base.startswith("Unnamed"):
            base = f"colonna {i + 1}"
        k = seen.get(base, 0)
        seen[base] = k + 1
        out.append(base if k == 0 else f"{base}.{k}")
    return out


def normalize(df: pd.DataFrame, m: Mapping) -> pd.DataFrame:
    """Tabella pulita: data, causale, entrate, uscite (positive), saldo_dichiarato."""
    df = df.loc[:, ~df.columns.duplicated()]  # nomi duplicati: vale la prima colonna
    dates = pd.to_datetime(df[m.data], errors="coerce", dayfirst=True)
    if dates.isna().any():
        raise ValueError(f"Date non valide alle righe: {[i + 1 for i in df.index[dates.isna()]]}")
    num = lambda col: pd.to_numeric(df[col], errors="coerce").fillna(0.0) if col else pd.Series(0.0, index=df.index)
    if m.importo:
        amt = num(m.importo)
        entrate, uscite = amt.clip(lower=0), (-amt).clip(lower=0)
    else:
        entrate, uscite = num(m.entrate), num(m.uscite)
    out = pd.DataFrame({
        "data": dates,
        "causale": df[m.causale].fillna("") if m.causale else "",
        "entrate": entrate, "uscite": uscite,
        "saldo_dichiarato": pd.to_numeric(df[m.saldo], errors="coerce") if m.saldo else float("nan")})
    return out.sort_values("data", kind="stable").reset_index(drop=True)


def saldo_mismatches(mov: pd.DataFrame) -> list[int]:
    """Righe (1-based) in cui il saldo dichiarato non torna con entrate-uscite cumulate."""
    running = (mov["entrate"] - mov["uscite"]).cumsum()
    diff = (mov["saldo_dichiarato"] - running).abs()
    return [i + 1 for i, d in enumerate(diff) if pd.notna(d) and d > SALDO_TOLERANCE]


def months_needed(mov: pd.DataFrame, opening: tuple[int, int]) -> list[tuple[int, int]]:
    """Mese del saldo iniziale + gennaio e dicembre dell'anno dell'ultimo movimento."""
    y = mov["data"].iloc[-1].year
    return sorted({opening, (y, 1), (y, 12)})


def build_workbook(mov: pd.DataFrame, rates: pd.DataFrame, iso: str, opening: tuple[int, int],
                   monthly: dict[tuple[int, int], float | None]) -> bytes:
    """Foglio 'Movimenti' (con formule) + 'Cambi' (giorno per giorno) + 'Medie mensili'."""
    missing = [k for k, v in monthly.items() if v is None]
    if missing:
        raise RatesError(f"Cambio medio mensile non ancora pubblicato per: {missing}")
    wb = Workbook()
    ws, wr, wm = wb.active, wb.create_sheet("Cambi"), wb.create_sheet("Medie mensili")
    ws.title = "Movimenti"
    bold = Font(bold=True)

    wr.append(["data", f"cambio {iso} per 1 EUR", "fonte"])
    for _, r in rates.iterrows():
        wr.append([r["data"].to_pydatetime(), r["cambio"], SOURCE])
    n_rates = wr.max_row

    wm.append(["anno", "mese", f"cambio medio {iso} per 1 EUR", "fonte"])
    keys = sorted(monthly)
    for y, mth in keys:
        wm.append([y, mth, monthly[(y, mth)], SOURCE_MONTHLY])
    opening_row = 2 + keys.index(opening)

    ws.append(["data", "causale", f"entrate ({iso})", f"uscite ({iso})", f"saldo ({iso})", "cambio",
               "data cambio usata", "entrate (EUR)", "uscite (EUR)", "saldo (EUR)", "gg", "saldo x gg (EUR)"])
    n = len(mov)
    lookup = lambda x, col: (f"=LOOKUP(A{x},Cambi!$A$2:$A${n_rates},Cambi!${col}$2:${col}${n_rates})")
    for i, r in mov.iterrows():
        x, first = i + 2, i == 0
        ws.append([r["data"].to_pydatetime(), r["causale"],
                   r["entrate"] or None, r["uscite"] or None,
                   f"=N(C{x})-N(D{x})" if first else f"=E{x-1}+N(C{x})-N(D{x})",
                   f"='Medie mensili'!C{opening_row}" if first else lookup(x, "B"),
                   "media mensile" if first else lookup(x, "A"),
                   f'=IF(C{x}="","",C{x}/F{x})', f'=IF(D{x}="","",D{x}/F{x})',
                   f"=N(H{x})-N(I{x})" if first else f"=J{x-1}+N(H{x})-N(I{x})",
                   f"=A{x+1}-A{x}" if i < n - 1 else f"=DATE(YEAR(A{x}),12,31)-A{x}",
                   f"=K{x}*J{x}"])
    for sheet in (ws, wr):
        for row in sheet.iter_rows(min_row=2, max_col=1):
            row[0].number_format = "DD/MM/YYYY"
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].number_format = "DD/MM/YYYY"
    for sheet in (ws, wr, wm):
        for col in range(1, sheet.max_column + 1):
            sheet.cell(row=1, column=col).font = bold
            sheet.column_dimensions[get_column_letter(col)].width = 18

    last = n + 1
    ws.cell(row=last + 2, column=2, value="giacenza media (EUR)").font = bold
    ws.cell(row=last + 2, column=10, value=f"=SUM(L2:L{last})/SUM(K2:K{last})")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def convert(mov: pd.DataFrame, iso: str, opening: tuple[int, int]) -> tuple[bytes, pd.DataFrame]:
    """Scarica i cambi necessari e costruisce l'xlsx. Ritorna (xlsx, tabella cambi)."""
    start = mov["data"].iloc[0].date() - dt.timedelta(days=LOOKBACK_DAYS)
    end = mov["data"].iloc[-1].date()
    rates = daily_rates(iso, start, end)
    monthly = {k: monthly_average(iso, *k) for k in months_needed(mov, opening)}
    return build_workbook(mov, rates, iso, opening, monthly), rates
